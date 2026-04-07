"""
Dilekçe şablonları endpoint'leri.
Şablonları listele, doldur, DOCX olarak indir.
"""

import io
import json
from datetime import datetime
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

router = APIRouter(tags=["templates"])

TEMPLATES_PATH = Path(__file__).parent.parent.parent / "templates" / "sablonlar.json"

# python-docx opsiyonel (DOCX çıktı için)
try:
    import docx
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def _load_templates() -> list[dict]:
    if not TEMPLATES_PATH.exists():
        return []
    data = json.loads(TEMPLATES_PATH.read_text(encoding="utf-8"))
    return data.get("sablonlar", [])


@router.get("/templates")
async def list_templates(hukuk_dali: Optional[str] = None):
    """Tüm dilekçe şablonlarını listele."""
    templates = _load_templates()
    if hukuk_dali:
        templates = [t for t in templates if t["hukuk_dali"] == hukuk_dali]
    return {
        "templates": [
            {
                "id": t["id"],
                "ad": t["ad"],
                "aciklama": t["aciklama"],
                "hukuk_dali": t["hukuk_dali"],
                "alanlar": t["alanlar"],
            }
            for t in templates
        ]
    }


@router.get("/templates/{template_id}")
async def get_template(template_id: str):
    """Tek bir şablonun detayını getir."""
    templates = _load_templates()
    for t in templates:
        if t["id"] == template_id:
            return t
    raise HTTPException(status_code=404, detail=f"Şablon bulunamadı: {template_id}")


class FillRequest(BaseModel):
    template_id: str
    alanlar: dict[str, str]
    format: str = "text"  # "text" veya "docx"


@router.post("/templates/fill")
async def fill_template(req: FillRequest):
    """Şablonu doldur, metin veya DOCX olarak döndür."""
    templates = _load_templates()
    template = None
    for t in templates:
        if t["id"] == req.template_id:
            template = t
            break
    if not template:
        raise HTTPException(status_code=404, detail=f"Şablon bulunamadı: {req.template_id}")

    # Zorunlu alanları kontrol et
    missing = []
    for alan in template["alanlar"]:
        if alan["zorunlu"] and alan["ad"] not in req.alanlar:
            # Varsayılan değer varsa onu kullan
            if alan.get("varsayilan"):
                req.alanlar[alan["ad"]] = alan["varsayilan"]
            else:
                missing.append(alan["etiket"])

    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Eksik zorunlu alanlar: {', '.join(missing)}",
        )

    # Tarih ekle
    req.alanlar.setdefault("tarih", datetime.now().strftime("%d/%m/%Y"))

    # Şablonu doldur
    filled = template["sablon"]
    for key, value in req.alanlar.items():
        filled = filled.replace(f"{{{key}}}", value)

    if req.format == "docx":
        if not HAS_DOCX:
            raise HTTPException(
                status_code=500,
                detail="DOCX desteği yok: pip install python-docx",
            )
        return _generate_docx(filled, template["ad"])

    return {
        "template_id": req.template_id,
        "ad": template["ad"],
        "text": filled,
    }


def _generate_docx(text: str, title: str) -> StreamingResponse:
    """Doldurulan şablondan DOCX oluştur."""
    doc = docx.Document()

    # Sayfa marjinleri
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # Metin satır satır ekle
    lines = text.split("\n")
    for line in lines:
        p = doc.add_paragraph()

        # Başlık satırları (tamamı büyük harf veya belirli formatlarda)
        stripped = line.strip()
        if stripped and (stripped.isupper() or stripped.startswith("SONUÇ VE TALEP")
                        or stripped.startswith("AÇIKLAMALAR") or stripped.startswith("KONU")):
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(12)
        elif stripped.startswith("I.") or stripped.startswith("II.") or stripped.startswith("III.") \
                or stripped.startswith("IV.") or stripped.startswith("V."):
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(11)
        else:
            run = p.add_run(line)
            run.font.size = Pt(11)

        run.font.name = "Times New Roman"
        p.paragraph_format.space_after = Pt(2)

    # Memory-only — diske yazma
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    filename = title.replace(" ", "_").lower() + ".docx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


class AIFillRequest(BaseModel):
    template_id: str
    description: str


@router.post("/templates/ai-fill")
async def ai_fill_template(req: AIFillRequest):
    """Dogal dil aciklamasindan dilekce alanlarini AI ile doldur."""
    from services.claude_service import generate_response

    templates = _load_templates()
    template = None
    for t in templates:
        if t["id"] == req.template_id:
            template = t
            break
    if not template:
        raise HTTPException(status_code=404, detail=f"Sablon bulunamadi: {req.template_id}")

    alan_lines = []
    for a in template["alanlar"]:
        z = "(ZORUNLU)" if a["zorunlu"] else "(opsiyonel)"
        alan_lines.append(f"- {a['ad']}: {a['etiket']} {z}")
    alan_listesi = "\n".join(alan_lines)

    prompt = (
        f"Asagidaki dilekce sablonu icin kullanicinin aciklamasindan alanlari doldur.\n\n"
        f"SABLON: {template['ad']}\n"
        f"ACIKLAMA: {req.description}\n\n"
        f"DOLDURULMASI GEREKEN ALANLAR:\n{alan_listesi}\n\n"
        f"Kullanicinin aciklamasindan cikarabildigin tum alanlari doldur.\n"
        f"Cikaramadigin alanlar icin makul varsayilan degerler kullan.\n"
        f"Yanitini SADECE JSON formatinda ver. Baska bir sey yazma."
    )

    try:
        response = generate_response(
            query=prompt,
            rag_context="",
            session_history=None,
            files_context="",
            model="claude-haiku-4-5-20251001",
        )

        text = response.strip()
        if text.startswith("```"):
            text = text.split("\n", 1)[1].rsplit("```", 1)[0].strip()

        alanlar = json.loads(text)
        alanlar.setdefault("tarih", datetime.now().strftime("%d/%m/%Y"))

        filled = template["sablon"]
        for key, value in alanlar.items():
            filled = filled.replace(f"{{{key}}}", str(value))

        return {
            "template_id": req.template_id,
            "ad": template["ad"],
            "alanlar": alanlar,
            "text": filled,
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"AI doldurma hatasi: {e}")
